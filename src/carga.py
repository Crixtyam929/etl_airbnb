"""Carga: se guardan las tablas limpias en una base SQLite nueva y en archivos XLSX."""

import logging
import sqlite3
from pathlib import Path

from openpyxl import load_workbook

CARPETA_DATOS = Path(__file__).resolve().parent.parent / "datos"
RUTA_BD = CARPETA_DATOS / "airbnb_bogota.db"
CARPETA_XLSX = CARPETA_DATOS / "xlsx"

# maximo filas en excel
MAXIMO_FILAS_EXCEL = 1_048_575

logger = logging.getLogger("carga")


class Carga:
    """Inserta las tablas en SQLite, las exporta a XLSX y verifica que quedaron completas."""

    def __init__(self, ruta_bd=RUTA_BD, carpeta_xlsx=CARPETA_XLSX):
        self.ruta_bd = ruta_bd
        self.carpeta_xlsx = carpeta_xlsx
        self.exportadas = []

    def cargar_sqlite(self, tablas):
        """Inserta cada tabla en SQLite; si ya existe se reemplaza para no duplicar datos."""
        conexion = sqlite3.connect(self.ruta_bd)
        try:
            for nombre, df in tablas.items():
                df.to_sql(nombre, conexion, if_exists="replace", index=False, chunksize=100_000)
                logger.info("Tabla '%s' insertada en SQLite: %d filas", nombre, len(df))
        finally:
            conexion.close()

    def exportar_xlsx(self, tablas):
        """Exporta cada tabla a su propio archivo XLSX, menos las que no caben en Excel."""
        self.carpeta_xlsx.mkdir(exist_ok=True)
        for nombre, df in tablas.items():
            if len(df) > MAXIMO_FILAS_EXCEL:
                logger.warning(
                    "La tabla '%s' tiene %d filas y no cabe en Excel: queda solo en SQLite", nombre, len(df)
                )
                continue
            ruta = self.carpeta_xlsx / f"{nombre}.xlsx"
            # Se evita que Excel convierta textos que empiezan por "=" o "http" en fórmulas o enlaces
            opciones = {"options": {"strings_to_formulas": False, "strings_to_urls": False}}
            df.to_excel(ruta, index=False, engine="xlsxwriter", engine_kwargs=opciones)
            self.exportadas.append(nombre)
            logger.info("Tabla '%s' exportada a %s: %d filas", nombre, ruta.name, len(df))

    def verificar(self, tablas):
        """Compara las filas de SQLite y de cada XLSX contra las filas transformadas."""
        conexion = sqlite3.connect(self.ruta_bd)
        try:
            for nombre, df in tablas.items():
                filas = conexion.execute(f"SELECT COUNT(*) FROM {nombre}").fetchone()[0]
                self.comparar(f"SQLite '{nombre}'", filas, len(df))
        finally:
            conexion.close()

        for nombre in self.exportadas:
            libro = load_workbook(self.carpeta_xlsx / f"{nombre}.xlsx", read_only=True)
            filas = libro.active.max_row - 1
            libro.close()
            self.comparar(f"XLSX '{nombre}'", filas, len(tablas[nombre]))

    @staticmethod
    def comparar(destino, cargadas, esperadas):
        if cargadas != esperadas:
            logger.error("%s quedó con %d filas y se esperaban %d", destino, cargadas, esperadas)
            raise ValueError(f"Carga incompleta en {destino}")
        logger.info("%s verificada: %d filas", destino, cargadas)

    def cargar(self, tablas):
        """Ejecuta la carga completa: SQLite, XLSX y verificación."""
        logger.info("Inicio de la carga en %s y %s", self.ruta_bd.name, self.carpeta_xlsx.name)
        self.cargar_sqlite(tablas)
        self.exportar_xlsx(tablas)
        self.verificar(tablas)
        logger.info("Carga terminada y verificada")
